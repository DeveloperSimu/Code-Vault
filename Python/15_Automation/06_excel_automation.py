from openpyxl import Workbook, load_workbook

file_name = "students.xlsx"

# Create Excel file
workbook = Workbook()
sheet = workbook.active

sheet.title = "Students"

sheet.append(["ID", "Name", "Age", "Course"])
sheet.append([1, "Srimanta", 21, "B.Sc IT"])
sheet.append([2, "Rahul", 22, "BCA"])
sheet.append([3, "Amit", 20, "B.Tech"])

workbook.save(file_name)

print("Excel file created successfully.")

# Read Excel file
workbook = load_workbook(file_name)
sheet = workbook["Students"]

print("\nStudent Records:\n")

for row in sheet.iter_rows(values_only=True):
    print(row)